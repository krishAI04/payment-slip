from __future__ import annotations

import io
import re
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook
import xlrd


class ValidationError(Exception):
    def __init__(self, message: str, rows: list[int] | None = None):
        self.message, self.rows = message, rows or []


@dataclass
class PaymentSlip:
    row: int
    work_order: str
    employee: str
    designation: str
    days: Decimal
    daily_rate: Decimal
    basic_wages: Decimal
    special_reward: Decimal
    total_pay: Decimal
    pf: Decimal
    esic: Decimal
    total_deduction: Decimal
    net_paid: Decimal
    wage_period: str
    establishment: str
    location: str
    nature_of_work: str

    def public(self):
        return {"workOrder": self.work_order, "employee": self.employee, "amount": money(self.net_paid), "row": self.row}


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def number(value, row: int, label: str, allow_na: bool = False) -> Decimal:
    if allow_na and clean(value).upper() in {"N/A", "NA", "-"}:
        return Decimal("0")
    if value is None or clean(value) == "":
        raise ValidationError(f"Missing {label} in Excel row {row}.", [row])
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"), ROUND_HALF_UP)
    except Exception:
        raise ValidationError(f"Invalid {label} in Excel row {row}.", [row])


def text_after(value, marker: str) -> str:
    return re.sub(re.escape(marker), "", clean(value), flags=re.I).strip(" :-")


def row_values(ws, row: int) -> list:
    return [ws.cell(row, column).value for column in range(1, 16)]


def first_text(values: list, contains: str) -> str:
    return next((clean(value) for value in values if contains.lower() in clean(value).lower()), "")


def parse_workbook(content: bytes) -> list[PaymentSlip]:
    try:
        if content.startswith(b"\xd0\xcf\x11\xe0"):
            workbook = xlrd.open_workbook(file_contents=content)
            worksheet = XlsSheet(workbook.sheet_by_index(0))
        else:
            workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
            worksheet = workbook.active
    except Exception:
        raise ValidationError("We could not read this Excel file. Please upload a valid .xlsx workbook.")
    starts = [row for row in range(1, worksheet.max_row + 1) if clean(worksheet.cell(row, 1).value).upper() == "FORM-XVII"]
    if not starts:
        raise ValidationError('This file does not contain the required "FORM-XVII" payment sections.')
    slips: list[PaymentSlip] = []
    problems: list[int] = []
    for index, start in enumerate(starts):
        end = starts[index + 1] if index + 1 < len(starts) else worksheet.max_row + 1
        header = [row_values(worksheet, row) for row in range(start, min(start + 6, end))]
        all_header_values = [value for row in header for value in row]
        work_order = text_after(first_text(all_header_values, "W. O. NO."), "W. O. NO.")
        establishment = text_after(first_text(all_header_values, "Name of Establishment"), "Name of Establishment")
        wage_period = text_after(first_text(all_header_values, "Wage Period"), "Wage Period")
        nature = next((clean(value) for value in all_header_values if "MAINT" in clean(value).upper() or "TEST" in clean(value).upper()), "")
        location = next((clean(value) for value in all_header_values if "BHEL" in clean(value).upper()), "")
        if not work_order:
            raise ValidationError(f"Missing Work Order Number near Excel row {start}.", [start])
        for row in range(start + 6, end):
            serial, employee = worksheet.cell(row, 1).value, clean(worksheet.cell(row, 2).value)
            if not isinstance(serial, int) or not employee:
                continue
            try:
                basic = number(worksheet.cell(row, 6).value, row, "Basic Wages")
                reward = number(worksheet.cell(row, 7).value or 0, row, "Special Reward")
                total = number(worksheet.cell(row, 8).value, row, "Total Pay")
                pf = number(worksheet.cell(row, 9).value or 0, row, "PF deduction", allow_na=True)
                esic = number(worksheet.cell(row, 10).value or 0, row, "ESIC deduction", allow_na=True)
                deduction = number(worksheet.cell(row, 11).value, row, "Total deduction")
                net = number(worksheet.cell(row, 12).value, row, "Net Amount Paid")
                slips.append(PaymentSlip(row, work_order, employee, clean(worksheet.cell(row, 3).value), number(worksheet.cell(row, 4).value, row, "days worked"), number(worksheet.cell(row, 5).value, row, "daily rate"), basic, reward, total, pf, esic, deduction, net, wage_period, establishment, location, nature))
            except ValidationError:
                problems.append(row)
    if problems:
        raise ValidationError("Some payment records contain missing or invalid required values.", sorted(set(problems)))
    if not slips:
        raise ValidationError("No employee payment records were found in the FORM-XVII sections.")
    return slips


class XlsSheet:
    """Tiny adapter so legacy .xls cells use the same 1-based access as openpyxl."""

    def __init__(self, sheet):
        self.sheet = sheet
        self.max_row = sheet.nrows

    def cell(self, row: int, column: int):
        return type("Cell", (), {"value": self.sheet.cell_value(row - 1, column - 1)})()


def money(value: Decimal) -> str:
    return f"{value:,.2f}".rstrip("0").rstrip(".")
